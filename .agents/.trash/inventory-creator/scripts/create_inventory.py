import argparse
import json
import subprocess
import sys
from collections import OrderedDict, defaultdict
from datetime import datetime
from decimal import Decimal
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import column_index_from_string, get_column_letter
from openpyxl.utils.datetime import from_excel
from openpyxl.worksheet.table import Table, TableStyleInfo


SCHEMA = "transco.sheet-ndjson"
SCHEMA_VERSION = 1
OUTPUT_SHEET = "Usage History"


def build_parser():
    parser = argparse.ArgumentParser(
        description=(
            "Create reusable NDJSON and a standardized monthly inventory-usage workbook."
        )
    )
    parser.add_argument("command", choices=["run"])
    parser.add_argument("input_xlsx", type=Path)
    parser.add_argument("output_xlsx", type=Path)
    parser.add_argument("--sheet", required=True)
    parser.add_argument("--header-row", type=int, required=True)
    parser.add_argument("--product-id-column", required=True)
    parser.add_argument("--product-name-column", required=True)
    parser.add_argument("--month-columns", nargs="+", required=True)
    parser.add_argument("--interim-dir", type=Path, required=True)
    return parser


def expand_columns(specifications):
    columns = []
    for specification in specifications:
        if ":" in specification:
            first, last = specification.split(":", 1)
            columns.extend(
                get_column_letter(index)
                for index in range(
                    column_index_from_string(first),
                    column_index_from_string(last) + 1,
                )
            )
        else:
            columns.append(specification.upper())
    return columns


def month_from_value(value):
    if isinstance(value, (int, float)):
        parsed = from_excel(value)
    else:
        parsed = pd.to_datetime(value).to_pydatetime()
    return datetime(parsed.year, parsed.month, 1)


def month_sequence(first_month, last_month):
    months = []
    current = first_month
    while current <= last_month:
        months.append(current)
        current = datetime(
            current.year + (current.month == 12),
            1 if current.month == 12 else current.month + 1,
            1,
        )
    return months


def normalized_identifier(value):
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def json_number(value):
    if value == value.to_integral_value():
        return int(value)
    return float(value)


def convert_source_to_ndjson(input_xlsx, output_ndjson):
    converter = (
        Path(__file__).resolve().parents[2]
        / "ndjson-for-sheet"
        / "scripts"
        / "sheet_to_ndjson.py"
    )
    subprocess.run(
        [sys.executable, str(converter), str(input_xlsx), str(output_ndjson)],
        check=True,
    )


def read_sheet_cells(source_ndjson, sheet_name):
    cells = {}
    with source_ndjson.open("r", encoding="utf-8") as input_file:
        for line in input_file:
            record = json.loads(line)
            if record["record_type"] == "cell" and record["sheet_name"] == sheet_name:
                cells[(record["row"], record["column"])] = record["value"]
    return cells


def build_inventory(cells, header_row, product_id_column, product_name_column, month_columns):
    product_id_index = column_index_from_string(product_id_column)
    product_name_index = column_index_from_string(product_name_column)
    month_indexes = [column_index_from_string(column) for column in month_columns]
    source_months = [
        (column, month_from_value(cells[(header_row, column)]))
        for column in month_indexes
    ]
    output_months = month_sequence(
        min(month for _, month in source_months),
        max(month for _, month in source_months),
    )

    products = OrderedDict()
    usage = defaultdict(Decimal)
    populated = set()
    source_rows = sorted({row for row, _ in cells if row > header_row})
    for row in source_rows:
        product_id = cells.get((row, product_id_index))
        product_name = cells.get((row, product_name_index))
        if product_id is None and product_name is None:
            continue

        product = (normalized_identifier(product_id), str(product_name))
        products.setdefault(product, None)
        for column, month in source_months:
            value = cells.get((row, column))
            if value is not None:
                usage[(product, month)] += Decimal(str(value))
                populated.add((product, month))

    return list(products), output_months, usage, populated


def cell_record(row, column, value, value_type, number_format):
    return {
        "record_type": "cell",
        "sheet_index": 0,
        "sheet_name": OUTPUT_SHEET,
        "address": f"{get_column_letter(column)}{row}",
        "row": row,
        "column": column,
        "value_type": value_type,
        "value": value,
        "formula": None,
        "number_format": number_format,
    }


def write_inventory_ndjson(path, source_name, products, months, usage, populated):
    last_column = 2 + len(months)
    last_row = 1 + len(products)
    with path.open("w", encoding="utf-8", newline="\n") as output_file:
        records = [
            {
                "record_type": "workbook",
                "schema": SCHEMA,
                "version": SCHEMA_VERSION,
                "source_name": source_name,
            },
            {
                "record_type": "sheet",
                "sheet_index": 0,
                "sheet_name": OUTPUT_SHEET,
                "dimension": f"A1:{get_column_letter(last_column)}{last_row}",
            },
            cell_record(1, 1, "Product ID", "string", "General"),
            cell_record(1, 2, "Product Name", "string", "General"),
        ]
        records.extend(
            cell_record(1, column, month.strftime("%b %Y"), "string", "General")
            for column, month in enumerate(months, start=3)
        )
        for record in records:
            output_file.write(json.dumps(record, ensure_ascii=False, separators=(",", ":")))
            output_file.write("\n")

        for row, product in enumerate(products, start=2):
            product_id, product_name = product
            product_records = [
                cell_record(row, 1, product_id, "string", "@"),
                cell_record(row, 2, product_name, "string", "General"),
            ]
            product_records.extend(
                cell_record(
                    row,
                    column,
                    json_number(usage[(product, month)]),
                    "number",
                    "#,##0.##",
                )
                for column, month in enumerate(months, start=3)
                if (product, month) in populated
            )
            for record in product_records:
                output_file.write(json.dumps(record, ensure_ascii=False, separators=(",", ":")))
                output_file.write("\n")


def write_workbook_from_ndjson(input_ndjson, output_xlsx):
    workbook = Workbook()
    worksheet = workbook.active
    dimension = None
    with input_ndjson.open("r", encoding="utf-8") as input_file:
        for line in input_file:
            record = json.loads(line)
            if record["record_type"] == "sheet":
                worksheet.title = record["sheet_name"]
                dimension = record["dimension"]
            elif record["record_type"] == "cell":
                cell = worksheet.cell(record["row"], record["column"], record["value"])
                cell.number_format = record["number_format"]

    header_fill = PatternFill("solid", fgColor="1F4E78")
    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = Font(color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    worksheet.row_dimensions[1].height = 26
    worksheet.column_dimensions["A"].width = 18
    worksheet.column_dimensions["B"].width = 32
    for column in range(3, worksheet.max_column + 1):
        worksheet.column_dimensions[get_column_letter(column)].width = 13
    worksheet.freeze_panes = "A2"
    worksheet.sheet_view.showGridLines = False

    table = Table(displayName="UsageHistoryTable", ref=dimension)
    table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=True,
        showColumnStripes=False,
    )
    worksheet.add_table(table)
    workbook.save(output_xlsx)


def run(args):
    args.interim_dir.mkdir(parents=True, exist_ok=True)
    args.output_xlsx.parent.mkdir(parents=True, exist_ok=True)
    source_ndjson = args.interim_dir / f"{args.input_xlsx.stem}.sheet.ndjson"
    inventory_ndjson = args.interim_dir / f"{args.output_xlsx.stem}.inventory.ndjson"

    convert_source_to_ndjson(args.input_xlsx, source_ndjson)
    cells = read_sheet_cells(source_ndjson, args.sheet)
    products, months, usage, populated = build_inventory(
        cells,
        args.header_row,
        args.product_id_column,
        args.product_name_column,
        expand_columns(args.month_columns),
    )
    write_inventory_ndjson(
        inventory_ndjson,
        args.output_xlsx.name,
        products,
        months,
        usage,
        populated,
    )
    write_workbook_from_ndjson(inventory_ndjson, args.output_xlsx)


def main():
    args = build_parser().parse_args()
    if args.command == "run":
        run(args)


if __name__ == "__main__":
    main()
