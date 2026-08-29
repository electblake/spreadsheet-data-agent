import json
from pathlib import Path

from loguru import logger
from openai import OpenAI
from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet
from SpreadsheetDataAgent.config import MODEL_ID, SYSTEM_RULES
from SpreadsheetDataAgent.helpers.document_files import select_file_from_name
from SpreadsheetDataAgent.helpers.embeddings import num_tokens
from SpreadsheetDataAgent.helpers.workbooks import (
    load_workbook_file,
    sheet_to_headers,
    word_list,
)

QUESTION_PROMPT = """Which worksheet in workbook {workbook_filename} tracks inventory/product usage over time?

Candidate worksheets: {workbook_sheetnames}

Workbook evidence:
{workbook_evidence}
"""


def suggest_inventory_sheet(
    wb: Workbook, workbook_filename: str
) -> tuple[Worksheet, int]:
    """Use OpenAI to predict which sheet tracks inventory usage over time."""

    logger.debug("Opening sheet prediction process")
    client = OpenAI()

    SYSTEM_PROMPT = (
        """You identify the worksheet that tracks inventory or product usage over time.

    Inventory/product usage means recurring product quantities organized by reporting period, such as month-by-month quantities shipped, consumed, issued, or used.

    Distinguish that historical usage pattern from line-item logistics data such as container numbers, ports, booking statuses, ready dates, and warehouse ETAs. Also distinguish authoritative business records from diagnostic outputs that report errors, mismatches, discrepancies, variances, exceptions, validation failures, or expected-versus-actual comparisons.

    TASK RULES
    - Prefer period-level product usage history over line-item shipment, container, booking, or warehouse logistics records.
    - Select a worksheet only when its rows directly record the business quantities used as the source of truth.
    - Reject audit, reconciliation, exception, and error-report worksheets. Products, dates, and quantities appearing only to describe a mismatch or other problem are not usage history.
    - If the evidence is insufficient or multiple worksheets satisfy the criteria equally, return UNDETERMINED.
    - Return only one candidate worksheet name with identical spelling, or UNDETERMINED.

    SYSTEM RULES
    - """
        + "\n- ".join(SYSTEM_RULES)
    )


    sheets_evidence = []
    for worksheet_name in wb.sheetnames:
        ws = wb[worksheet_name]

        worksheet_headers = sheet_to_headers(ws)
        worksheet_header_words = [
            (
                row_reference,
                word_list(
                    [
                        value
                        for value in header_values
                        if isinstance(value, str)
                    ]
                ),
            )
            for row_reference, header_values in worksheet_headers
        ]

        sheets_evidence.append(f"""
            ---
            Worksheet Name: '{worksheet_name}'
            Worksheet Headers (words): {worksheet_header_words}
        """)

    workbook_evidence = ("""Workbook Filename: {workbook_filename}

    Worksheet Evidence List:
    """ + "\n".join(sheets_evidence))

    workbook_sheetnames=", ".join(wb.sheetnames)

    logger.debug(f"workbook_sheetnames: {workbook_sheetnames}")
    logger.debug(f"workbook_filename: {workbook_filename}")
    logger.debug(f"workbook_evidence (tokens): {num_tokens(workbook_evidence)}")

    logger.debug("Asking openai ({})\nQuestion: \n ```\n{}\n```", MODEL_ID, QUESTION_PROMPT)
    response = client.responses.create(
        model=MODEL_ID,
        instructions=SYSTEM_PROMPT,
        input=QUESTION_PROMPT.format(
            workbook_filename=workbook_filename,
            workbook_sheetnames=workbook_sheetnames,
            workbook_evidence=workbook_evidence,
        ),
        text={
            "format": {
                "type": "json_schema",
                "name": "sheet_prediction",
                "strict": True,
                "schema": {
                    "type": "object",
                    "properties": {
                        "sheet_name": {
                            "type": "string",
                            "enum": [*wb.sheetnames, "UNDETERMINED"],
                        }
                    },
                    "required": ["sheet_name"],
                    "additionalProperties": False,
                },
            }
        },
    )

    prediction = json.loads(response.output_text)
    logger.debug("Prediction response ({}): {}", MODEL_ID, prediction)

    if prediction["sheet_name"] == "UNDETERMINED":
        raise ValueError("Could not determine the inventory usage worksheet")

    inventory_sheet = wb[prediction["sheet_name"]]
    sheet_index = wb.sheetnames.index(prediction["sheet_name"])
    return inventory_sheet, sheet_index

def main(file: str):
    logger.debug("Opening workbook load process")
    if file_path := select_file_from_name(file):
        if wb := load_workbook_file(file_path):
            inventory_sheet, sheet_index = suggest_inventory_sheet(wb, file_path) # pyright: ignore[reportArgumentType]
            logger.debug(
                "Suggested Sheet: '{}' (index: {}, rows: {})",
                inventory_sheet.title,
                sheet_index,
                inventory_sheet.max_row,
            )
            return inventory_sheet.title, inventory_sheet, sheet_index
        else:
            logger.critical("could not load_workbook_file with {}", file_path)
    else:
        logger.critical("select_file_from_name returned none for {}", file)

    return None

# handle cli-like callers
if __name__ == "__main__":
    import argparse

    parser = argparse.ArgumentParser()
    parser.add_argument(
        "-f",
        "--file",
        required=True,
        type=Path,
        help="Workbook path or filename.",
    )
    args = parser.parse_args()

    main(args.file)

# TODO: define fastmcp-compatible interface below
