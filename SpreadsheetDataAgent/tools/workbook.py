from pathlib import Path

from openpyxl import Workbook


def predict_layout_options(wb: Workbook):
    return {
        "sheetnames": wb.sheetnames,

    }

def load_workbook(file_path: str | Path) -> Workbook | None:
    """returns read-only file handler for xlsx/xlsm files"""
    file_path = Path(file_path)
    from openpyxl import load_workbook
    if file_path and file_path.exists and file_path.is_file():
        # https://openpyxl.readthedocs.io/en/stable/tutorial.html#loading-from-a-file
        return load_workbook(filename=file_path, read_only=True, keep_vba=True, rich_text=True, keep_links=True, data_only=False)
    return None

def read_defined_name_ranges(wb: Workbook) -> list[tuple[str, str, list[str]]]:
    defined_names = list(wb.defined_names.items())
    for sheet in wb.worksheets:
        defined_names.extend(sheet.defined_names.items())

    ranges_by_name_and_sheet = {}
    for name, defined_name in defined_names:
        for sheet_name, cell_range in defined_name.destinations:
            ranges_by_name_and_sheet.setdefault((name, sheet_name), []).append(
                cell_range.replace("$", "")
            )

    return [
        (name, sheet_name, cell_ranges)
        for (name, sheet_name), cell_ranges in ranges_by_name_and_sheet.items()
    ]
