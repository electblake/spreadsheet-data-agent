import json
from pathlib import Path

from loguru import logger
from openai import OpenAI
from openpyxl import Workbook

from SpreadsheetDataAgent.config import MODEL_ID, SYSTEM_RULES
from SpreadsheetDataAgent.helpers.workbooks import load_workbook, search_workbooks

QUESTION_PROMPT = """Which worksheet in workbook {workbook_filename} tracks inventory/product usage over time?

Candidate worksheets: {workbook_sheetnames}

Worksheet evidence:
{sheet_evidence}
"""

def predict_inventory_sheet_name(wb: Workbook, file: str | Path):
    """Use OpenAI to predict which sheet tracks inventory usage over time."""

    logger.debug("Opening sheet prediction process")
    client = OpenAI()

    SYSTEM_PROMPT = (
        """You identify the worksheet that tracks inventory or product usage over time.

    Inventory/product usage means recurring product quantities organized by reporting period, such as month-by-month quantities shipped, consumed, issued, or used.

    Distinguish that historical usage pattern from line-item logistics data such as container numbers, ports, booking statuses, ready dates, and warehouse ETAs. Also distinguish authoritative business records from diagnostic outputs that report errors, mismatches, discrepancies, variances, exceptions, validation failures, or expected-versus-actual comparisons.

    TASK RULES
    - Prefer period-level product usage history over line-item shipment, container, booking, or warehouse logistics records.
    - Select a worksheet only when its rows directly record the business quantities used as the source of truth.
    - Reject audit, reconciliation, exception, and error-report worksheets. Products, dates, and quantities appearing only to describe a mismatch or other problem are not usage history.
    - If the evidence is insufficient or multiple worksheets satisfy the criteria equally, return UNDETERMINED.
    - Return only one candidate worksheet name with identical spelling, or UNDETERMINED.

    SYSTEM RULES
    - """
        + "\n- ".join(SYSTEM_RULES)
    )
    sheet_evidence = "\n\n".join(
        "\n".join(
            [
                f"Sheet: {ws.title}",
                f"Dimensions: {ws.max_row} rows x {ws.max_column} columns",
                *[
                    ", ".join(
                        f"{cell.coordinate}={cell.value!r}"
                        for cell in row
                        if cell.value is not None
                    )
                    for row in ws.iter_rows(
                        min_row=1,
                        max_row=min(ws.max_row, 15),
                        max_col=min(ws.max_column, 20),
                    )
                    if any(cell.value is not None for cell in row)
                ],
            ]
        )
        for ws in wb.worksheets
    )

    file_path = Path(file)
    workbook_filename = file_path.name
    workbook_sheetnames=", ".join(wb.sheetnames)

    logger.debug(f"Creating openai inference ({MODEL_ID})")
    logger.debug(f"- workbook_sheetnames: {workbook_sheetnames}")
    logger.debug(f"- workbook_filename: {workbook_filename}")

    response = client.responses.create(
        model=MODEL_ID,
        instructions=SYSTEM_PROMPT,
        input=QUESTION_PROMPT.format(
            workbook_filename=workbook_filename,
            workbook_sheetnames=workbook_sheetnames,
            sheet_evidence=sheet_evidence,
        ),
        text={
            "format": {
                "type": "json_schema",
                "name": "sheet_prediction",
                "strict": True,
                "schema": {
                    "type": "object",
                    "properties": {
                        "sheet_name": {
                            "type": "string",
                            "enum": [*wb.sheetnames, "UNDETERMINED"],
                        }
                    },
                    "required": ["sheet_name"],
                    "additionalProperties": False,
                },
            }
        },
    )
    logger.debug("Closed openai inference ({MODEL_ID})")

    prediction = json.loads(response.output_text)
    logger.debug("Prediction response: {}", prediction)

    if prediction["sheet_name"] == "UNDETERMINED":
        raise ValueError("Could not determine the inventory usage worksheet")

    inventory_sheet = wb[prediction["sheet_name"]]
    logger.debug(
        "Inventory sheet: '{}' (rows: {})",
        inventory_sheet.title,
        inventory_sheet.max_row,
    )
    logger.debug("Defined names available:")
    for name, defined_name in wb.defined_names.items():
        destinations = [
            f"'{sheet_name}'!{cell_range}"
            for sheet_name, cell_range in defined_name.destinations
        ]
        logger.debug("- {} ({})", name, ", ".join(destinations))
    logger.debug("Closed sheet prediction process")


if __name__ == "__main__":
    import argparse

    logger.debug("Starting argument parsing")
    parser = argparse.ArgumentParser()
    parser.add_argument(
        "-f",
        "--file",
        required=True,
        type=Path,
        help="Workbook path or filename.",
    )
    args = parser.parse_args()

    file_path = search_workbooks(args.file)
    logger.debug("Opening workbook load process")
    wb = load_workbook(file_path)
    logger.debug("Closed workbook load process")

    if wb:
        predict_inventory_sheet_name(wb, file_path)
    else:
        raise FileNotFoundError(file_path)
