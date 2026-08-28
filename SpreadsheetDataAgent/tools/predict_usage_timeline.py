import json

from loguru import logger
from openai import OpenAI
from openpyxl import Workbook
from workbook import load_workbook, read_defined_name_ranges

from SpreadsheetDataAgent.config import MODEL_ID, SYSTEM_RULES

QUESTION_PROMPT = """Identify the product rows and month usage columns in the supplied Excel {named_range} named-range data.

Worksheet evidence:
{sheet_evidence}
"""

NAMED_RANGE = "USAGE_TIMELINE"

def predict_usage_timeline(wb: Workbook):
    """TBD"""

    client = OpenAI()

    SYSTEM_PROMPT = (
        """You identify the product rows and month usage that tracks inventory/product usage over time.

    Inventory/product usage means recurring product quantities organized by reporting period, such as month-by-month quantities shipped, consumed, issued, or used.

    TASK RULES:
    - Distinguish that historical usage pattern from line-item logistics data such as container numbers, ports, booking statuses, ready dates, and warehouse ETAs.
    - A positive match requires evidence of time periods, identifiable products, and recurring quantity values.

    SYSTEM RULES:
    - """
        + "\n- ".join(SYSTEM_RULES)
    )



    defined_name_ranges = read_defined_name_ranges(wb)
    usage_timeline_ranges = [
        (sheet_name, cell_range)
        for name, sheet_name, cell_ranges in defined_name_ranges
        if name == NAMED_RANGE
        for cell_range in cell_ranges
    ]
    usage_timeline = [
        {
            "sheet_name": sheet_name,
            "cell_range": cell_range,
            "rows": [
                {cell.coordinate: cell.value for cell in row}
                for row in wb[sheet_name][cell_range]
            ],
        }
        for sheet_name, cell_range in usage_timeline_ranges
    ]

    logger.debug("usage_timeline: {}", usage_timeline)

    sheet_evidence = json.dumps(usage_timeline)

    response = client.responses.create(
        model=MODEL_ID,
        instructions=SYSTEM_PROMPT,
        input=QUESTION_PROMPT.format(
            workbook_sheetnames=", ".join(wb.sheetnames),
            sheet_evidence=sheet_evidence,
            named_range=NAMED_RANGE,
        ),
        text={
            "format": {
                "type": "json_schema",
                "name": "usage_timeline",
                "strict": True,
                "schema": {
                    "type": "object",
                    "properties": {
                        "product_rows": {
                            "type": "array",
                            "items": {
                                "type": "object",
                                "properties": {
                                    "sheet_name": {"type": "string"},
                                    "cell_range": {"type": "string"},
                                    "values": {
                                        "type": "array",
                                        "items": {
                                            "type": [
                                                "string",
                                                "number",
                                                "boolean",
                                                "null",
                                            ]
                                        },
                                    },
                                },
                                "required": [
                                    "sheet_name",
                                    "cell_range",
                                    "values",
                                ],
                                "additionalProperties": False,
                            },
                        },
                        "month_usage_columns": {
                            "type": "array",
                            "items": {
                                "type": "object",
                                "properties": {
                                    "sheet_name": {"type": "string"},
                                    "cell_range": {"type": "string"},
                                    "values": {
                                        "type": "array",
                                        "items": {
                                            "type": [
                                                "string",
                                                "number",
                                                "boolean",
                                                "null",
                                            ]
                                        },
                                    },
                                },
                                "required": [
                                    "sheet_name",
                                    "cell_range",
                                    "values",
                                ],
                                "additionalProperties": False,
                            },
                        }
                    },
                    "required": ["product_rows", "month_usage_columns"],
                    "additionalProperties": False,
                },
            }
        },
    )

    return json.loads(response.output_text)

if __name__ == "__main__":
    import argparse

    parser = argparse.ArgumentParser()
    parser.add_argument("-f", "--file", required=True)
    args = parser.parse_args()

    logger.debug("Opening workbook: {}", args.file)
    if wb := load_workbook(args.file):
        logger.debug("Opened workbook: {}", args.file)
        predict_usage_timeline(wb)
        logger.debug("Closing workbook: {}", args.file)
        wb.close()
        logger.debug("Closed workbook: {}", args.file)
    else:
        raise FileNotFoundError(args.file)
