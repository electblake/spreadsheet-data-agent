from pathlib import Path

from openpyxl import load_workbook


def build_workbook_overview(file_path: str | Path) -> str:
    workbook = load_workbook(file_path, read_only=True, data_only=False)
    worksheet_count = len(workbook.worksheets)
    lines = [
        f"Workbook overview ({Path(file_path).name}):",
        f"- Worksheet names (case-sensitive): {', '.join(workbook.sheetnames)}",
        f"- Active worksheet: {workbook.active.title}",
        f"- Worksheet dimensions ({worksheet_count}/{worksheet_count} shown):",
    ]

    for index, worksheet in enumerate(workbook.worksheets, start=1):
        lines.append(
            f"  {index}) {worksheet.title}: {worksheet.calculate_dimension()} "
            f"(max_row={worksheet.max_row}, max_col={worksheet.max_column})"
        )

    workbook.close()
    return "\n".join(lines)
