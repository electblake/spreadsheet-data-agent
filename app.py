import json
from pathlib import Path
import tomllib

import gradio as gr
import pandas as pd
import spaces
from huggingface_hub import hf_hub_download
from llama_cpp import Llama
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter


PROJECT_VERSION = tomllib.loads(
    Path(__file__).with_name("pyproject.toml").read_text(encoding="utf-8")
)["project"]["version"]
MODEL_REPO = "mradermacher/Spreadsheet-RL-4B-GGUF"
XLSX_PREVIEW_ROWS = 100
XLSX_PREVIEW_COLUMNS = 50
QUANT_FILES = {
    "Q2_K · 1.9 GB": "Spreadsheet-RL-4B.Q2_K.gguf",
    "Q3_K_S · 2.2 GB": "Spreadsheet-RL-4B.Q3_K_S.gguf",
    "Q3_K_M · 2.3 GB · lower quality": "Spreadsheet-RL-4B.Q3_K_M.gguf",
    "Q3_K_L · 2.5 GB": "Spreadsheet-RL-4B.Q3_K_L.gguf",
    "IQ4_XS · 2.6 GB": "Spreadsheet-RL-4B.IQ4_XS.gguf",
    "Q4_K_S · 2.7 GB · recommended": "Spreadsheet-RL-4B.Q4_K_S.gguf",
    "Q4_K_M · 2.8 GB · recommended": "Spreadsheet-RL-4B.Q4_K_M.gguf",
    "Q5_K_S · 3.2 GB": "Spreadsheet-RL-4B.Q5_K_S.gguf",
    "Q5_K_M · 3.3 GB": "Spreadsheet-RL-4B.Q5_K_M.gguf",
    "Q6_K · 3.7 GB · very good quality": "Spreadsheet-RL-4B.Q6_K.gguf",
    "Q8_0 · 4.8 GB · best quality": "Spreadsheet-RL-4B.Q8_0.gguf",
    "f16 · 8.9 GB": "Spreadsheet-RL-4B.f16.gguf",
}

model = None
active_quant = None


def build_workbook_overview(file_path: str | Path) -> pd.DataFrame:
    workbook = load_workbook(file_path, read_only=True, data_only=False)
    active_sheet = workbook.active.title
    overview = pd.DataFrame(
        [
            {
                "Sheet": worksheet.title,
                "Active": worksheet.title == active_sheet,
                "Used range": worksheet.calculate_dimension(force=True),
                "Rows": worksheet.max_row,
                "Columns": worksheet.max_column,
            }
            for worksheet in workbook.worksheets
        ]
    )
    workbook.close()
    return overview


def preview_xlsx_sheet(file_path: str | Path, sheet_name: str) -> pd.DataFrame:
    workbook = load_workbook(file_path, read_only=True, data_only=False)
    worksheet = workbook[sheet_name]
    worksheet.calculate_dimension(force=True)
    column_count = min(worksheet.max_column, XLSX_PREVIEW_COLUMNS)
    row_count = min(worksheet.max_row, XLSX_PREVIEW_ROWS)
    preview = pd.DataFrame(
        worksheet.iter_rows(
            min_row=1,
            max_row=row_count,
            min_col=1,
            max_col=column_count,
            values_only=True,
        ),
        columns=[get_column_letter(index) for index in range(1, column_count + 1)],
    )
    workbook.close()
    return preview


def load_xlsx_workflow(
    file_path: str | Path,
) -> tuple[pd.DataFrame, gr.Dropdown, pd.DataFrame]:
    overview = build_workbook_overview(file_path)
    sheet_names = overview["Sheet"].tolist()
    selected_sheet = sheet_names[0]
    return (
        overview,
        gr.Dropdown(choices=sheet_names, value=selected_sheet),
        preview_xlsx_sheet(file_path, selected_sheet),
    )


def download_quant(quantization: str) -> None:
    hf_hub_download(repo_id=MODEL_REPO, filename=QUANT_FILES[quantization])


def file_to_text(file_path: str | None) -> str:
    if file_path is None:
        return ""

    path = Path(file_path)
    suffix = path.suffix.lower()

    if suffix in {".xlsx", ".xls"}:
        sheets = pd.read_excel(path, sheet_name=None)
        return "\n\n".join(
            f"## Sheet: {sheet_name}\n{frame.to_csv(index=False)}"
            for sheet_name, frame in sheets.items()
        )

    if suffix == ".csv":
        return pd.read_csv(path).to_csv(index=False)

    if suffix == ".tsv":
        return pd.read_csv(path, sep="\t").to_csv(index=False)

    return path.read_text(encoding="utf-8")


@spaces.GPU(duration=120)
def generate(
    system_prompt: str,
    user_prompt: str,
    attachment: str | None,
    quantization: str,
) -> str:
    global active_quant, model

    quant_file = QUANT_FILES[quantization]
    if active_quant != quantization:
        model = None
        active_quant = None
        model = Llama(
            model_path=hf_hub_download(repo_id=MODEL_REPO, filename=quant_file),
            n_ctx=4096,
            n_gpu_layers=-1,
            verbose=True,
        )
        active_quant = quantization

    attachment_text = file_to_text(attachment)
    user_content = user_prompt
    if attachment_text:
        user_content = f"{user_prompt}\n\n<attachment>\n{attachment_text}\n</attachment>"

    messages = [
        {
            "role": "system",
            "content": (
                f"{system_prompt}\n\nAfter private reasoning, answer once with only the "
                "requested deliverable. Follow the user's output format exactly. Do not "
                "restate analysis, reasoning, or self-correction in the final answer."
            ),
        },
        {"role": "user", "content": user_content},
    ]
    completion = model.create_chat_completion(
        messages=messages,
        max_tokens=512,
        temperature=0.6,
        top_p=0.95,
        top_k=20,
    )
    return completion["choices"][0]["message"]["content"].rsplit("</think>", 1)[-1].strip()


@spaces.GPU(duration=60)
def generate_xlsx_table(
    system_prompt: str,
    user_prompt: str,
    attachment: str,
    sheet_name: str,
    quantization: str,
) -> pd.DataFrame:
    global active_quant, model

    quant_file = QUANT_FILES[quantization]
    if active_quant != quantization:
        model = None
        active_quant = None
        model = Llama(
            model_path=hf_hub_download(repo_id=MODEL_REPO, filename=quant_file),
            n_ctx=4096,
            n_gpu_layers=-1,
            verbose=True,
        )
        active_quant = quantization

    overview = build_workbook_overview(attachment).to_csv(index=False)
    preview = preview_xlsx_sheet(attachment, sheet_name).to_csv(index=False)
    messages = [
        {
            "role": "system",
            "content": (
                f"{system_prompt}\n\nReturn only a JSON table object with a columns "
                "array and a rows array. Every row must contain one string value per "
                "column."
            ),
        },
        {
            "role": "user",
            "content": (
                f"Task:\n{user_prompt}\n\nWorkbook overview:\n{overview}\n"
                f"Selected sheet preview ({sheet_name}, first {XLSX_PREVIEW_ROWS} rows "
                f"and {XLSX_PREVIEW_COLUMNS} columns):\n{preview}"
            ),
        },
    ]
    completion = model.create_chat_completion(
        messages=messages,
        response_format={
            "type": "json_object",
            "schema": {
                "type": "object",
                "properties": {
                    "columns": {
                        "type": "array",
                        "items": {"type": "string"},
                    },
                    "rows": {
                        "type": "array",
                        "items": {
                            "type": "array",
                            "items": {"type": "string"},
                        },
                    },
                },
                "required": ["columns", "rows"],
                "additionalProperties": False,
            },
        },
        max_tokens=1024,
        temperature=0.6,
        top_p=0.95,
        top_k=20,
    )
    response = json.loads(completion["choices"][0]["message"]["content"])
    return pd.DataFrame(response["rows"], columns=response["columns"])


CSS = """
.gradio-container {
    width: min(calc(100% - 32px), 1600px) !important;
    max-width: 1600px !important;
    margin-inline: auto !important;
}
.agent-panel { border: 2px dashed #79b5ce; border-radius: 18px; padding: 8px; }
.output-panel { border: 2px dashed #f0aeb7; border-radius: 18px; padding: 8px; }
"""

with gr.Blocks(css=CSS, title="Spreadsheet Data Agent") as demo:
    gr.Markdown(
        f"""
        # Spreadsheet Data Agent · v{PROJECT_VERSION}

        [Code](https://github.com/electblake/Spreadsheet-RL-Data-Agent) | [Demo](https://huggingface.co/spaces/electblake/spreadsheet-data-agent) | [Paper](https://arxiv.org/abs/2605.22642) | [Spreadsheet-RL Model](https://huggingface.co/Spreadsheet-RL/Spreadsheet-RL-4B)

        Send instructions and optional file context to Spreadsheet-RL-4B. This first
        inference surface implements the prompt-and-file entry point from the agent diagram.
        """
    )

    with gr.Tabs(selected="basic-data"):
        with gr.Tab("Basic data only", id="basic-data"):
            gr.Markdown(
                "Uses the original inference workflow: uploaded files are converted to "
                "plain data context and sent to the model with the prompt."
            )

            with gr.Row():
                with gr.Column(scale=1, elem_classes="agent-panel"):
                    gr.Markdown("### RL data input")
                    system_prompt = gr.Textbox(
                        label="System prompt",
                        value=(
                            "You are a spreadsheet reasoning assistant. Inspect the supplied "
                            "spreadsheet or text context and answer the user's request precisely."
                        ),
                        lines=5,
                    )
                    user_prompt = gr.Textbox(
                        label="User prompt",
                        placeholder="Describe the spreadsheet task or ask a question…",
                        lines=8,
                    )
                    attachment = gr.File(
                        label="Optional file context",
                        file_types=[".txt", ".md", ".json", ".csv", ".tsv", ".xlsx", ".xls"],
                        type="filepath",
                    )
                    quantization = gr.Dropdown(
                        choices=list(QUANT_FILES),
                        value="Q4_K_M · 2.8 GB · recommended",
                        label="Spreadsheet-RL-4B quantization",
                        info="Static GGUF quants published by mradermacher; Q4_K_M is the reference recommendation.",
                    )
                    run = gr.Button("Run inference", variant="primary")

                with gr.Column(scale=1, elem_classes="output-panel"):
                    gr.Markdown("### Agent response")
                    response = gr.Textbox(
                        label="Generated text",
                        lines=28,
                        buttons=["copy"],
                    )

        with gr.Tab("XLSX workflow", id="xlsx-workflow"):
            gr.Markdown(
                """
                ## XLSX workbook workflow

                Upload an Excel workbook to inspect its sheets and preview its table data
                before running table-focused inference.
                """
            )

            xlsx_attachment = gr.File(
                label="XLSX workbook",
                file_types=[".xlsx"],
                type="filepath",
            )
            workbook_overview = gr.Dataframe(
                headers=["Sheet", "Active", "Used range", "Rows", "Columns"],
                datatype=["str", "bool", "str", "number", "number"],
                label="Workbook summary",
                interactive=False,
                buttons=["fullscreen", "copy"],
                show_search="filter",
            )
            xlsx_sheet = gr.Dropdown(
                label="Preview sheet",
                choices=[],
            )
            xlsx_preview = gr.Dataframe(
                label=(
                    f"Selected sheet preview (first {XLSX_PREVIEW_ROWS} rows and "
                    f"{XLSX_PREVIEW_COLUMNS} columns)"
                ),
                interactive=False,
                max_height=520,
                buttons=["fullscreen", "copy"],
                show_row_numbers=True,
                show_search="filter",
            )

            with gr.Row():
                with gr.Column(scale=1, elem_classes="agent-panel"):
                    gr.Markdown("### Table inference input")
                    xlsx_system_prompt = gr.Textbox(
                        label="System prompt",
                        value=(
                            "You are a spreadsheet data assistant. Analyze the workbook "
                            "summary and selected sheet preview, then return the requested "
                            "result as a table."
                        ),
                        lines=5,
                    )
                    xlsx_user_prompt = gr.Textbox(
                        label="User prompt",
                        placeholder="Describe the table to derive from this workbook…",
                        lines=8,
                    )
                    xlsx_quantization = gr.Dropdown(
                        choices=list(QUANT_FILES),
                        value="Q4_K_M · 2.8 GB · recommended",
                        label="Spreadsheet-RL-4B quantization",
                        info="Static GGUF quants published by mradermacher; Q4_K_M is the reference recommendation.",
                    )
                    xlsx_run = gr.Button("Run table inference", variant="primary")

                with gr.Column(scale=1, elem_classes="output-panel"):
                    gr.Markdown("### Table response")
                    xlsx_response = gr.Dataframe(
                        label="Generated table",
                        interactive=False,
                        max_height=720,
                        buttons=["fullscreen", "copy"],
                        show_row_numbers=True,
                        show_search="filter",
                    )

    gr.Markdown(
        """
        ---

        ### Citation

        If you use Spreadsheet-RL-4B, please cite the model's paper:

        ```bibtex
        @misc{chi2026spreadsheetrl,
          title         = {Spreadsheet-RL: Advancing Large Language Model Agents on Realistic Spreadsheet Tasks via Reinforcement Learning},
          author        = {Banghao Chi and Yining Xie and Mingyuan Wu and Jingcheng Yang and Jize Jiang and Zhaoheng Li and Shengyi Qian and Minjia Zhang and Klara Nahrstedt and Rui Hou and Xiangjun Fan and Hanchao Yu},
          year          = {2026},
          eprint        = {2605.22642},
          archivePrefix = {arXiv},
          primaryClass  = {cs.AI},
          doi           = {10.48550/arXiv.2605.22642},
          url           = {https://arxiv.org/abs/2605.22642}
        }
        ```

        Citation from the [Spreadsheet-RL-4B model card](https://huggingface.co/Spreadsheet-RL/Spreadsheet-RL-4B#citation).
        """
    )

    run.click(
        fn=download_quant,
        inputs=quantization,
        outputs=None,
        show_progress="full",
    ).then(
        fn=generate,
        inputs=[system_prompt, user_prompt, attachment, quantization],
        outputs=response,
        api_name="generate",
        show_progress="full",
    )

    xlsx_attachment.upload(
        fn=load_xlsx_workflow,
        inputs=xlsx_attachment,
        outputs=[workbook_overview, xlsx_sheet, xlsx_preview],
        show_progress="full",
    )
    xlsx_sheet.change(
        fn=preview_xlsx_sheet,
        inputs=[xlsx_attachment, xlsx_sheet],
        outputs=xlsx_preview,
        show_progress="full",
    )
    xlsx_run.click(
        fn=download_quant,
        inputs=xlsx_quantization,
        outputs=None,
        show_progress="full",
    ).then(
        fn=generate_xlsx_table,
        inputs=[
            xlsx_system_prompt,
            xlsx_user_prompt,
            xlsx_attachment,
            xlsx_sheet,
            xlsx_quantization,
        ],
        outputs=xlsx_response,
        api_name="generate_xlsx_table",
        show_progress="full",
    )

demo.queue().launch(mcp_server=True)
