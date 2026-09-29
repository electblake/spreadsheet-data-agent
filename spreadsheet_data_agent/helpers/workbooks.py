import csv
import json
from io import StringIO

from openpyxl import Workbook
from openpyxl.utils.cell import get_column_letter, quote_sheetname, range_boundaries
from openpyxl.worksheet.worksheet import Worksheet


def read_defined_name_ranges(wb: Workbook) -> list[tuple[str, str, list[str]]]:
    defined_names = list(wb.defined_names.items())
    for sheet in wb.worksheets:
        defined_names.extend(sheet.defined_names.items())

    ranges_by_name_and_sheet = {}
    for name, defined_name in defined_names:
        for sheet_name, cell_range in defined_name.destinations:
            ranges_by_name_and_sheet.setdefault((name, sheet_name), []).append(
                cell_range.replace("$", "")
            )

    return [
        (name, sheet_name, cell_ranges)
        for (name, sheet_name), cell_ranges in ranges_by_name_and_sheet.items()
    ]

def named_range_to_data(wb: Workbook, named_range:str) -> list[list[int|str|float]]:
    named_range_data = []
    for sheet_name, cell_range in wb.defined_names[named_range].destinations:
        min_col, min_row, max_col, max_row = range_boundaries(cell_range)
        named_range_data.extend(
            [
                value
                if type(value) in (str, int, float)
                else ""
                if value is None
                else str(value)
                for value in row
            ]
            for row in wb[sheet_name].iter_rows(
                min_row=min_row,
                max_row=max_row,
                min_col=min_col,
                max_col=max_col,
                values_only=True,
            )
        )
    return named_range_data

def sheet_to_data(ws: Worksheet) -> list:
    wb_data=[]
    wb_data.extend(
        [
            value
            if type(value) in (str, int, float)
            else ""
            if value is None
            else str(value)
            for value in row
        ]
        for row in ws.iter_rows(values_only=True)
    )
    return wb_data

def to_data(wb: Workbook) -> list[list[str|int|float]]:
    wb_data = []
    for ws in wb.worksheets:
        wb_data.append({
            "worksheet_title": ws.title,
            "worksheet_max_rows": ws.max_row,
            "worksheet_max_column": ws.max_column,
            "named_ranges": read_defined_name_ranges(wb)
        })
        wb_data.extend(sheet_to_data(ws))
    return wb_data

def to_text(wb: Workbook) -> str:
    wb_data = to_data(wb)
    return json.dumps(wb_data, ensure_ascii=False)

def to_csv(wb: Workbook) -> str:
    return "\n".join(sheet_to_csv(sheet) for sheet in wb.worksheets)

def sheet_to_csv(sheet: Worksheet) -> str:
    output = StringIO(newline="")
    writer = csv.writer(output)
    writer.writerows(sheet.values)
    return output.getvalue()

def sheet_to_headers(sheet: Worksheet) -> list[tuple[str, tuple[object, ...]]]:
    """returns only rows that __could__ be headers/column names"""
    header_rows = []
    for row_number, row in enumerate(sheet.iter_rows(), start=1):
        non_empty_cells = [cell for cell in row if cell.value is not None]
        if not non_empty_cells:
            continue

        first_non_empty_cell = non_empty_cells[0]
        if (
            first_non_empty_cell.data_type == "f"
            or not isinstance(first_non_empty_cell.value, str)
        ):
            continue

        non_formula_cells = [
            cell for cell in non_empty_cells if cell.data_type != "f"
        ]
        text_cells = [
            cell for cell in non_formula_cells if isinstance(cell.value, str)
        ]

        if len(text_cells) / len(non_formula_cells) < 0.75:
            continue

        row_reference = (
            f"{quote_sheetname(sheet.title)}!"
            f"A{row_number}:{get_column_letter(len(row))}{row_number}"
        )

        header_rows.append((row_reference, tuple(cell.value for cell in row)))

    return header_rows


def word_list(input: str | list[str]) -> list[str]:
    """make list of unique words from split string(s)"""
    words = set()
    if not isinstance(input, list):
        input = [input]
    for text in input:
        words.update(text.split())
    return sorted(words)
